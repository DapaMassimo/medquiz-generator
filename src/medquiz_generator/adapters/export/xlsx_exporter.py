"""XLSX export (openpyxl): one row per question, readable by non-technical reviewers."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from medquiz_generator.domain.errors import ExportError
from medquiz_generator.domain.models import QuizItem

HEADERS = [
    "id",
    "materia",
    "argomento",
    "difficoltà",
    "domanda",
    "A",
    "B",
    "C",
    "D",
    "E",
    "risposta corretta",
    "spiegazione",
    "sezione",
    "percorso sezione",
    "pagina",
    "pagine sezione",
    "citazione",
    "stato",
    "motivo",
    "modello",
    "run id",
]

WIDTHS = {
    "domanda": 70,
    "spiegazione": 70,
    "citazione": 60,
    "percorso sezione": 40,
    "sezione": 30,
    "motivo": 40,
}


def row_of(item: QuizItem) -> list[object]:
    texts = {o.id: o.text for o in item.options}
    return [
        item.id,
        item.materia,
        item.argomento,
        item.difficulty.value,
        item.question,
        texts.get("A", ""),
        texts.get("B", ""),
        texts.get("C", ""),
        texts.get("D", ""),
        texts.get("E", ""),
        item.correct_option_id,
        item.explanation,
        item.source.section_title,
        item.source.section_path,
        item.source.page if item.source.page is not None else "",
        f"{item.source.page_start}-{item.source.page_end}",
        item.source.quote,
        item.status.status.value,
        item.status.reason,
        item.model_metadata.model,
        item.model_metadata.run_id,
    ]


class XlsxExporter:
    @property
    def format_name(self) -> str:
        return "xlsx"

    def export(self, items: list[QuizItem], target: Path) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        if sheet is None:  # pragma: no cover - openpyxl always creates one sheet
            raise ExportError("Workbook has no active sheet")
        sheet.title = "Quesiti"
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for item in items:
            sheet.append(row_of(item))
        for index, header in enumerate(HEADERS, start=1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = WIDTHS.get(header, 16)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(target)
        except OSError as exc:
            raise ExportError(f"Cannot write {target}: {exc}") from exc
        return target
