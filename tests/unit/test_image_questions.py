from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from medquiz_generator.adapters.export.xlsx_exporter import XlsxExporter
from medquiz_generator.application import artifacts
from medquiz_generator.application.classify import ClassifyStage
from medquiz_generator.application.context import RunContext
from medquiz_generator.application.generate import GenerateStage
from medquiz_generator.application.plan import PlanStage
from medquiz_generator.application.verify import VerifyStage
from medquiz_generator.domain.models import (
    FigureCatalog,
    FigureRef,
    ImageRef,
    QuestionStatus,
    StructuredDocument,
)
from tests.fakes.fake_llm import (
    FakeGenerator,
    FakeImageQuestionGenerator,
    FakeJudge,
    FakeTaxonomyBuilder,
)
from tests.unit.test_figures import PNG
from tests.unit.test_verify import good_item


def _prepare(ctx: RunContext, document: StructuredDocument) -> None:
    run_dir = ctx.store.run_dir(ctx.run_id)
    (run_dir / "assets" / "figures").mkdir(parents=True)
    (run_dir / "assets" / "figures" / "f001.png").write_bytes(PNG)
    figures = [
        FigureRef(
            id="f001",
            page=2,
            section_id="s003",
            caption="Figura 1 ECG a 12 derivazioni",
            path="assets/figures/f001.png",
            kind="clinical_image",
            question_worthy=True,
        ),
        FigureRef(
            id="f002",
            page=3,
            section_id="s004",
            caption="Figura 2 Schema",
            path="assets/figures/f002.png",
            kind="diagram",
        ),
    ]
    ctx.store.save(ctx.run_id, artifacts.DOCUMENT, document.model_copy(update={"figures": figures}))
    ctx.store.save(ctx.run_id, artifacts.FIGURES, FigureCatalog(run_id=ctx.run_id, figures=figures))
    PlanStage(min_unit_chars=400).run(ctx, 4)
    ClassifyStage(FakeTaxonomyBuilder()).run(ctx)


def test_generate_adds_one_question_per_question_worthy_figure(
    ctx: RunContext, fixture_document: StructuredDocument
) -> None:
    _prepare(ctx, fixture_document)
    image_gen = FakeImageQuestionGenerator(ctx.tracker)
    items = GenerateStage(FakeGenerator(), [], image_generator=image_gen).run(ctx)
    assert image_gen.calls == ["f001"]  # the diagram is never used
    image_items = [i for i in items if i.image is not None]
    assert len(image_items) == 1 and len(items) == 5
    item = image_items[0]
    assert item.id == "fig-f001-q01" and item.model_metadata.unit_id == "fig-f001"
    assert item.image == ImageRef(
        figure_id="f001",
        path="assets/figures/f001.png",
        caption="Figura 1 ECG a 12 derivazioni",
        page=2,
    )
    assert item.source.section_id == "s003" and item.materia
    # Resumable: a second run does not regenerate the figure question.
    GenerateStage(FakeGenerator(), [], image_generator=image_gen).run(ctx)
    assert image_gen.calls == ["f001"]
    usage = [r for r in ctx.tracker.records() if r.model == "fake-vision-generator"]
    assert usage and usage[0].item_id == "fig-f001"


def test_verify_flags_image_questions_and_shows_the_image_to_the_judge(
    ctx: RunContext, fixture_document: StructuredDocument
) -> None:
    _prepare(ctx, fixture_document)
    items = GenerateStage(FakeGenerator(), [], image_generator=FakeImageQuestionGenerator()).run(
        ctx
    )
    judge = FakeJudge()
    verified = VerifyStage(judge).run(ctx)
    by_id = {i.id: i for i in verified}
    image_item = by_id["fig-f001-q01"]
    assert image_item.status.status == QuestionStatus.FLAGGED
    assert image_item.status.reason.startswith("image question: human review required")
    assert judge.images_seen == ["fig-f001-q01"]
    assert all(i.status.status == QuestionStatus.VERIFIED for i in verified if i.image is None)
    # Without the review requirement the judge's pass stands.
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_GENERATED, list(items))
    relaxed = VerifyStage(FakeJudge(), image_questions_require_review=False)
    assert relaxed.verify_item(image_item, fixture_document, PNG).final_status == "verified"


def test_xlsx_export_carries_the_image_path(tmp_path: Path) -> None:
    item = good_item("fig-x").model_copy(
        update={"image": ImageRef(figure_id="f001", path="assets/figures/f001.png", caption="c")}
    )
    path = XlsxExporter().export([item, good_item("plain")], tmp_path / "quiz.xlsx")
    rows = list(load_workbook(path)["Quesiti"].iter_rows(values_only=True))
    assert rows[0][-1] == "immagine"
    assert rows[1][-1] == "assets/figures/f001.png" and rows[2][-1] in ("", None)


def test_report_leaves_image_questions_out_of_delivered_styles(ctx: RunContext) -> None:
    from medquiz_generator.application.report import build_run_report

    plain = good_item("u001-q01").with_status(QuestionStatus.VERIFIED, "ok")
    on_figure = (
        good_item("fig-f001-q01")
        .model_copy(update={"image": ImageRef(figure_id="f001", path="a/f001.png", caption="c")})
        .with_status(QuestionStatus.FLAGGED, "image question: human review required")
    )
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_VERIFIED, [plain, on_figure])
    report = build_run_report(ctx)
    assert report.image_questions == 1
    assert report.styles_delivered == {plain.style.value: 1}
