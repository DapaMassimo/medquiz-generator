from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from medquiz_generator.adapters.export.json_exporter import JsonExporter
from medquiz_generator.adapters.export.xlsx_exporter import HEADERS, XlsxExporter
from medquiz_generator.adapters.storage.file_run_store import FileGoldenStore
from medquiz_generator.application import artifacts
from medquiz_generator.application.context import RunContext
from medquiz_generator.application.eval import EvalStage, averages_of, compare
from medquiz_generator.application.export import ExportStage, select_deliverable
from medquiz_generator.application.report import ReportStage, build_run_report
from medquiz_generator.application.review import ReviewService
from medquiz_generator.domain.errors import MedQuizError, StageInputMissingError
from medquiz_generator.domain.models import (
    EvalItemResult,
    EvalReport,
    QuestionStatus,
    QuizItem,
    RunReport,
    StructuredDocument,
    StyleScores,
    UsageRecord,
)
from tests.fakes.fake_llm import FakeStyleJudge
from tests.unit.test_verify import GOOD_QUOTE, good_item


def sample_items() -> list[QuizItem]:
    return [
        good_item("q1").with_status(QuestionStatus.VERIFIED, "ok"),
        good_item("q2").with_status(QuestionStatus.FLAGGED, "judge: dubbio"),
        good_item("q3").with_status(QuestionStatus.DISCARDED, "deterministic: quote_missing"),
    ]


def test_select_deliverable() -> None:
    items = sample_items()
    assert [i.id for i in select_deliverable(items, include_flagged=False)] == ["q1"]
    assert [i.id for i in select_deliverable(items, include_flagged=True)] == ["q1", "q2"]


def test_json_and_xlsx_exporters(tmp_path: Path) -> None:
    items = sample_items()[:2]
    json_path = JsonExporter().export(items, tmp_path / "out" / "quiz.json")
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert [p["id"] for p in payload] == ["q1", "q2"]
    assert payload[0]["source"]["quote"] == items[0].source.quote
    xlsx_path = XlsxExporter().export(items, tmp_path / "out" / "quiz.xlsx")
    sheet = load_workbook(xlsx_path)["Quesiti"]
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert rows[1][0] == "q1" and rows[1][11] == "A" and rows[2][18] == "flagged"
    assert rows[0][4] == "stile" and rows[1][4] == "vignetta"
    assert len(rows) == 3


def test_export_stage_requires_verified_items(ctx: RunContext) -> None:
    with pytest.raises(StageInputMissingError):
        ExportStage([JsonExporter()]).run(ctx)


def test_export_stage_writes_files_and_meta(ctx: RunContext) -> None:
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_VERIFIED, list(sample_items()))
    written = ExportStage([JsonExporter(), XlsxExporter()]).run(ctx)
    assert set(written) == {"json", "xlsx"}
    assert all(p.exists() for p in written.values())
    assert ctx.meta().exported_count == 1
    written = ExportStage([JsonExporter()]).run(ctx, include_flagged=True)
    payload = json.loads(written["json"].read_text(encoding="utf-8"))
    assert [p["id"] for p in payload] == ["q1", "q2"]
    assert ctx.meta().exported_count == 2


def test_export_approved_only_uses_review_state(ctx: RunContext, tmp_path: Path) -> None:
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_DEDUPED, list(sample_items()))
    review = ReviewService(ctx.store, FileGoldenStore(tmp_path / "golden"))
    review.decide(ctx.run_id, "q2", "approved", "buona")
    review.decide(ctx.run_id, "q3", "approved", "non dovrebbe passare: scartata")
    written = ExportStage([JsonExporter()]).run(ctx, approved_only=True)
    payload = json.loads(written["json"].read_text(encoding="utf-8"))
    assert [p["id"] for p in payload] == ["q2"]
    assert written["json"].name == "quiz.approved.json"


def test_review_service_decisions_and_promotion(ctx: RunContext, tmp_path: Path) -> None:
    golden = FileGoldenStore(tmp_path / "golden")
    review = ReviewService(ctx.store, golden)
    assert review.items(ctx.run_id) == []
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_VERIFIED, list(sample_items()))
    with pytest.raises(MedQuizError):
        review.promote(ctx.run_id, "q1")  # not approved yet
    review.decide(ctx.run_id, "q1", "approved", "ottima")
    review.decide(ctx.run_id, "q3", "approved")
    assert [i.id for i in review.approved_items(ctx.run_id)] == ["q1", "q3"]
    review.promote(ctx.run_id, "q1")
    assert [i.id for i in golden.load_all()] == ["q1"]
    with pytest.raises(MedQuizError):
        review.promote(ctx.run_id, "q3")  # discarded items can never be promoted
    with pytest.raises(MedQuizError):
        review.promote(ctx.run_id, "missing")
    review.clear(ctx.run_id, "q1")
    assert "q1" not in review.state(ctx.run_id).decisions


def test_report_aggregates_usage_and_counts(ctx: RunContext) -> None:
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_DEDUPED, list(sample_items()))
    ctx.tracker.set_stage("generate")
    ctx.tracker.set_context(unit_id="u001")
    ctx.tracker.record(
        UsageRecord(
            run_id="", stage="", model="fake-generator", input_tokens=1_000_000, output_tokens=0
        )
    )
    ctx.tracker.set_stage("verify")
    ctx.tracker.set_context(unit_id="u001", item_id="q1")
    ctx.tracker.record(
        UsageRecord(run_id="", stage="", model="fake-judge", input_tokens=0, output_tokens=500_000)
    )
    report = ReportStage().run(ctx)
    assert (report.generated, report.verified, report.flagged, report.discarded) == (3, 1, 1, 1)
    assert report.usage_total.cost_usd == pytest.approx(2.0)
    assert report.usage_by_stage["generate"].input_tokens == 1_000_000
    assert report.usage_by_unit["u001"].calls == 2
    assert report.usage_by_item["q1"].cost_usd == pytest.approx(1.0)
    assert report.cost_per_delivered_question_usd == pytest.approx(1.0)
    assert report.models == {"generate": "fake-generator", "verify": "fake-judge"}
    assert report.discard_reasons == {"deterministic: quote_missing": 1}
    assert report.schema_version == 1
    stored = ctx.store.load(ctx.run_id, artifacts.RUN_REPORT, RunReport)
    assert stored.run_id == ctx.run_id
    assert build_run_report(ctx).units_covered == 1


def test_eval_stage_scores_verified_items_and_compare(
    ctx: RunContext, tmp_path: Path, fixture_document: StructuredDocument
) -> None:
    ctx.store.save(ctx.run_id, artifacts.DOCUMENT, fixture_document)
    ctx.store.save_list(ctx.run_id, artifacts.QUESTIONS_DEDUPED, list(sample_items()))
    golden = FileGoldenStore(tmp_path / "golden")
    golden.add("old-run", good_item("gold"))
    judge = FakeStyleJudge()
    report = EvalStage(judge, [], golden).run(ctx)
    assert judge.calls == ["q1"]
    # The style judge receives the context window around the quote, not the bare quote.
    assert "camice bianco" in judge.passages[0] and len(judge.passages[0]) > len(GOOD_QUOTE)
    assert report.golden_regression_count == 1
    assert report.averages["groundedness"] == 5.0 and report.averages["mean"] == 4.5
    assert ctx.store.exists(ctx.run_id, artifacts.EVAL_REPORT)
    with_flagged = EvalStage(judge, [], golden, include_flagged=True).run(ctx, limit=1)
    assert len(with_flagged.items) == 1

    other = EvalReport(
        run_id="other",
        judge_model="fake",
        golden_style_count=0,
        golden_regression_count=0,
        items=[
            EvalItemResult(
                item_id="x",
                scores=StyleScores(
                    ssm_style=2,
                    distractor_quality=2,
                    explanation_clarity=2,
                    groundedness=2,
                    rationale="",
                ),
            )
        ],
        averages=averages_of([]),
    )
    other = other.model_copy(update={"averages": averages_of(other.items)})
    comparison = compare(other, report)
    assert comparison.delta["groundedness"] == 3.0
    assert comparison.delta["mean"] == 2.5
    assert averages_of([])["mean"] == 0.0
